import csv

from docx import Document

from russian_vocab_desktop.export import export_entries
from russian_vocab_desktop.models import LemmaEntry


def test_export_entries_writes_separate_pos_files(tmp_path):
    output = tmp_path / "lemmas.csv"
    files = export_entries(
        [
            LemmaEntry("книга", 4, "noun"),
            LemmaEntry("читать", 3, "verb"),
        ],
        output,
        export_mode="separate_by_pos",
        include_frequency=True,
        include_pos=True,
        include_header=True,
        sort_mode="frequency_desc",
    )

    assert {path.name for path in files} == {"lemmas_nouns.csv", "lemmas_verbs.csv"}

    rows = list(csv.reader((tmp_path / "lemmas_nouns.csv").open(encoding="utf-8")))
    assert rows == [["lemma", "frequency", "pos"], ["книга", "4", "noun"]]


def test_export_entries_writes_docx_output(tmp_path):
    output = tmp_path / "lemmas.docx"
    files = export_entries(
        [
            LemmaEntry("книга", 4, "noun"),
            LemmaEntry("читать", 3, "verb"),
        ],
        output,
        export_mode="single_file",
        include_frequency=True,
        include_pos=True,
        include_header=True,
        sort_mode="frequency_desc",
    )

    assert files == (output,)

    document = Document(output)
    assert document.paragraphs[0].text == "Russian Vocabulary Export"

    table = document.tables[0]
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    assert rows == [
        ["lemma", "frequency", "pos"],
        ["книга", "4", "noun"],
        ["читать", "3", "verb"],
    ]
