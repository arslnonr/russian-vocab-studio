"""Export helpers for CSV and Word outputs."""

from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path
from typing import Callable

from .lemmatize import cyrillic_sort_key
from .models import ExportMode, LemmaEntry, POS_LABELS, SortMode


def export_entries(
    entries: list[LemmaEntry],
    output_path: Path,
    *,
    export_mode: ExportMode,
    include_frequency: bool,
    include_pos: bool,
    include_header: bool,
    sort_mode: SortMode,
) -> tuple[Path, ...]:
    output_path, write_entries = _resolve_output_writer(output_path)
    sorted_entries = _sort_entries(entries, sort_mode)

    if export_mode == "single_file":
        return (
            write_entries(
                output_path,
                sorted_entries,
                include_frequency,
                include_pos,
                include_header,
            ),
        )

    grouped: dict[str, list[LemmaEntry]] = defaultdict(list)
    for entry in sorted_entries:
        grouped[entry.pos].append(entry)

    written: list[Path] = []
    if export_mode == "both":
        written.append(
            write_entries(
                output_path,
                sorted_entries,
                include_frequency,
                include_pos,
                include_header,
            )
        )

    stem = output_path.stem
    directory = output_path.parent
    suffix = output_path.suffix
    for pos_group, entries_for_pos in grouped.items():
        label = POS_LABELS[pos_group]
        path = directory / f"{stem}_{label}{suffix}"
        written.append(
            write_entries(path, entries_for_pos, include_frequency, include_pos, include_header)
        )
    return tuple(written)


def _sort_entries(entries: list[LemmaEntry], sort_mode: SortMode) -> list[LemmaEntry]:
    if sort_mode == "alphabetical":
        return sorted(entries, key=lambda item: cyrillic_sort_key(item.lemma))
    return sorted(entries, key=lambda item: (-item.frequency, cyrillic_sort_key(item.lemma)))


def _write_csv(
    path: Path,
    entries: list[LemmaEntry],
    include_frequency: bool,
    include_pos: bool,
    include_header: bool,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = ["lemma"]
    if include_frequency:
        headers.append("frequency")
    if include_pos:
        headers.append("pos")

    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        if include_header:
            writer.writerow(headers)
        for entry in entries:
            row = [entry.lemma]
            if include_frequency:
                row.append(entry.frequency)
            if include_pos:
                row.append(entry.pos)
            writer.writerow(row)
    return path


def _write_docx(
    path: Path,
    entries: list[LemmaEntry],
    include_frequency: bool,
    include_pos: bool,
    include_header: bool,
) -> Path:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is not installed. Install project dependencies before exporting Word documents."
        ) from exc

    path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    document.add_heading("Russian Vocabulary Export", level=1)

    columns = ["lemma"]
    if include_frequency:
        columns.append("frequency")
    if include_pos:
        columns.append("pos")

    if not entries and not include_header:
        document.add_paragraph("No lemmas matched the current filters.")
        document.save(path)
        return path

    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"

    start_index = 0
    if include_header:
        for index, header in enumerate(columns):
            table.rows[0].cells[index].text = header
    elif entries:
        _fill_docx_row(
            table.rows[0].cells,
            entries[0],
            include_frequency=include_frequency,
            include_pos=include_pos,
        )
        start_index = 1

    for entry in entries[start_index:]:
        _fill_docx_row(
            table.add_row().cells,
            entry,
            include_frequency=include_frequency,
            include_pos=include_pos,
        )

    document.save(path)
    return path


def _fill_docx_row(
    cells,
    entry: LemmaEntry,
    *,
    include_frequency: bool,
    include_pos: bool,
) -> None:
    values = [entry.lemma]
    if include_frequency:
        values.append(str(entry.frequency))
    if include_pos:
        values.append(entry.pos)

    for index, value in enumerate(values):
        cells[index].text = value


def _resolve_output_writer(path: Path) -> tuple[Path, Callable[..., Path]]:
    suffix = path.suffix.lower()
    if not suffix:
        normalized = path.with_suffix(".csv")
        return normalized, _write_csv
    if suffix == ".csv":
        return path.with_suffix(".csv"), _write_csv
    if suffix == ".docx":
        return path.with_suffix(".docx"), _write_docx
    raise ValueError("Unsupported output format. Choose a .csv or .docx file.")
